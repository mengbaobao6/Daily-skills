#!/usr/bin/env python3
"""Offline and fake-API tests for the high-throughput Excel workflow."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook

from clone_engine import build, extract_render_xml
from excel_workflow import (
    ADD_CONCURRENCY,
    Workflow,
    WorkbookLedger,
    compile_excel_rows,
)


SOURCE_XML = ""


def list_response(rows: list[dict] | None = None) -> dict:
    return {
        "alibaba_icbu_product_list_response": {
            "products": {"alibaba_product_brief_response": rows or []}
        }
    }


def product_get_response(source_id: str = "1601000000000") -> dict:
    return {
        "alibaba_icbu_product_get_response": {
            "product": {
                "product_id": source_id,
                "category_id": 201273078,
                "gmt_modified": "2026-07-27 12:00:00",
                "status": "approved",
                "display": "Y",
            }
        }
    }


def render_response(xml: str) -> dict:
    return {
        "alibaba_icbu_product_schema_render_response": {
            "biz_success": True,
            "data": xml,
        }
    }


class FakeReadClient:
    def __init__(self):
        self.lock = threading.Lock()
        self.calls: dict[str, int] = {}

    def request(self, method: str, business: dict | None = None, timeout: int = 60) -> dict:
        with self.lock:
            self.calls[method] = self.calls.get(method, 0) + 1
        if method == "alibaba.icbu.product.get":
            return product_get_response()
        if method == "alibaba.icbu.product.schema.render":
            return render_response(SOURCE_XML)
        if method == "alibaba.icbu.product.schema.get":
            request = business["param_product_top_publish_request"]
            if request.get("cat_id") != 201273078:
                raise AssertionError("schema.get must use cat_id")
            return {
                "alibaba_icbu_product_schema_get_response": {
                    "biz_success": True,
                    "data": SOURCE_XML,
                }
            }
        if method == "alibaba.icbu.product.list":
            return list_response()
        raise AssertionError(method)


class FakeAddClient:
    def __init__(self, ambiguous_title: str = "", render_available: bool = True,
                 inventory_available: bool = True, inventory_timeout: bool = False):
        self.ambiguous_title = ambiguous_title
        self.render_available = render_available
        self.inventory_available = inventory_available
        self.inventory_timeout = inventory_timeout
        self.lock = threading.Lock()
        self.active = 0
        self.max_active = 0
        self.counter = 0
        self.rendered: dict[str, str] = {}
        self.inventory: dict[str, dict[tuple[str, str], int]] = {}
        self.inventory_update_calls = 0

    def request(self, method: str, business: dict | None = None, timeout: int = 60) -> dict:
        if method == "alibaba.icbu.product.schema.add":
            xml = business["param_product_top_publish_request"]["xml"]
            if self.ambiguous_title and self.ambiguous_title in xml:
                raise TimeoutError("simulated ambiguous timeout")
            with self.lock:
                self.active += 1
                self.max_active = max(self.max_active, self.active)
                self.counter += 1
                sequence = self.counter
                product_id = str(1700000000000 + sequence)
            root = ET.fromstring(xml)
            inventory: dict[tuple[str, str], int] = {}
            for index, row in enumerate(root.findall("./field[@id='sku']/complex-values"), start=1):
                sku_id = str(900000 + sequence * 100 + index)
                sku_field = row.find("field[@id='skuId']")
                if sku_field is None:
                    sku_field = ET.SubElement(row, "field", {"id": "skuId", "type": "input"})
                for old in list(sku_field.findall("value")):
                    sku_field.remove(old)
                ET.SubElement(sku_field, "value").text = sku_id
                stock = row.find("field[@id='skuStock']/values/value")
                warehouse = stock.get("warehouseCode") if stock is not None else "CN_LOCAL_01"
                inventory[(sku_id, warehouse)] = 0
            with self.lock:
                self.rendered[product_id] = ET.tostring(root, encoding="unicode")
                self.inventory[product_id] = inventory
            time.sleep(0.05)
            with self.lock:
                self.active -= 1
            return {
                "alibaba_icbu_product_schema_add_response": {
                    "biz_success": True,
                    "productId": product_id,
                    "request_id": f"request-{product_id}",
                }
            }
        if method == "alibaba.icbu.product.schema.render":
            product_id = str(business["param_product_top_publish_request"]["product_id"])
            if self.render_available:
                return render_response(self.rendered[product_id])
            return {"alibaba_icbu_product_schema_render_response": {"biz_success": False}}
        if method == "alibaba.icbu.product.sku.inventory.get":
            product_id = str(business["product_id"])
            if not self.inventory_available:
                return {"alibaba_icbu_product_sku_inventory_get_response": {"result": {"data_list": []}}}
            rows = [{
                "sku_id": sku_id,
                "inventory_code": warehouse,
                "inventory": value,
            } for (sku_id, warehouse), value in self.inventory[product_id].items()]
            return {
                "alibaba_icbu_product_sku_inventory_get_response": {
                    "result": {"data_list": rows}
                }
            }
        if method == "alibaba.icbu.product.inventory.update":
            request = business["request_param"]
            product_id = str(request["product_id"])
            with self.lock:
                self.inventory_update_calls += 1
                if self.inventory_timeout:
                    raise TimeoutError("simulated inventory ambiguity")
                for change in request["inventory_list"]:
                    pair = (str(change["sku_id"]), str(change["inventory_code"]))
                    delta = int(change["inventory"])
                    self.inventory[product_id][pair] += delta if change["operate"] == "plus" else -delta
            return {"alibaba_icbu_product_inventory_update_response": {"result": {"success": True}}}
        raise AssertionError(method)


class FakeVerifyClient:
    def __init__(self, product_id: str, title: str, render_xml: str, inventory: int):
        self.product_id = product_id
        self.title = title
        self.render_xml = render_xml
        self.inventory = inventory

    def request(self, method: str, business: dict | None = None, timeout: int = 60) -> dict:
        if method == "alibaba.icbu.product.get":
            return {
                "alibaba_icbu_product_get_response": {
                    "product": {
                        "product_id": self.product_id,
                        "category_id": 201273078,
                        "subject": self.title,
                        "status": "approved",
                        "display": "Y",
                    }
                }
            }
        if method == "alibaba.icbu.product.schema.render":
            return render_response(self.render_xml)
        if method == "alibaba.icbu.product.sku.inventory.get":
            return {
                "alibaba_icbu_product_sku_inventory_get_response": {
                    "result": {
                        "data_list": [{
                            "sku_id": "900001",
                            "inventory_code": "CN_LOCAL_01",
                            "inventory": self.inventory,
                        }]
                    }
                }
            }
        raise AssertionError(method)


def create_workbook(root: Path, row_count: int = 1) -> Path:
    images = root / "images"
    images.mkdir()
    (images / "1.jpg").write_bytes(b"same-image")
    plan = {
        "variant_axes": [{"name": "Color", "values": ["Gray"]}],
        "sku_defaults": {"stock_target": 10, "warehouse_code": "CN_LOCAL_01"},
        "price_tiers": [{"quantity": 10, "price": "10.50"}],
        "moq": 10,
        "lead_times": [{"quantity": 10, "day": 15}],
        "package": {"length": 40, "width": 15, "height": 15, "weight": 1.5},
    }
    (root / "plan.json").write_text(json.dumps(plan), encoding="utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "产品发布清单"
    sheet.append(["任务ID", "源商品ID", "新产品标题", "新产品图片路径", "新产品规格参数", "SKU方案文件路径"])
    for index in range(row_count):
        sheet.append([
            f"P{index + 1:03d}",
            "1601000000000",
            f"Unique Product Title {index + 1}",
            "images",
            "Material=PP",
            "plan.json",
        ])
    notes = workbook.create_sheet("填写说明")
    notes.append(["字段", "说明"])
    notes.append(["源商品ID", "必填"])
    path = root / "products.xlsx"
    workbook.save(path)
    return path


def ready_product(key: str, title: str) -> dict:
    return {
        "item_key": key,
        "source_product_id": "1601000000000",
        "category_id": 201273078,
        "title": title,
        "inventory_mode": "deferred",
        "main_images": [{"file_id": "123", "url": "https://sc04.alicdn.com/kf/Htest.jpg"}],
        "detail_galleries": [{"gallery": "200", "display_name": "Scene", "images": [
            {"url": "https://sc04.alicdn.com/kf/Htest.jpg"},
        ]}],
        "category_attribute_overrides": {"Material": "PP"},
        "variant_axes": [{"name": "Color", "values": ["Gray"]}],
        "sku_defaults": {"stock_target": 10, "warehouse_code": "CN_LOCAL_01"},
        "price_tiers": [{"quantity": 10, "price": "10.50"}],
        "moq": 10,
        "lead_times": [{"quantity": 10, "day": 15}],
        "package": {"length": 40, "width": 15, "height": 15, "weight": 1.5},
    }


class ExcelWorkflowTests(unittest.TestCase):
    def test_instruction_sheet_is_ignored(self):
        with tempfile.TemporaryDirectory() as folder:
            path = create_workbook(Path(folder))
            rows = compile_excel_rows(path)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["sheet"], "产品发布清单")

    def test_twenty_rows_fetch_one_source_get_and_render(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 20)
            client = FakeReadClient()
            workflow = Workflow(path, root / "artifacts", client=client)
            summary = workflow.prepare(upload_images=False)
            self.assertEqual(summary["unique_sources"], 1)
            self.assertEqual(client.calls["alibaba.icbu.product.get"], 1)
            self.assertEqual(client.calls["alibaba.icbu.product.schema.render"], 1)
            self.assertEqual(client.calls["alibaba.icbu.product.list"], 1)

    def test_images_upload_once_with_five_way_limit(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 2)
            images = root / "images"
            for index in range(2, 7):
                (images / f"{index}.jpg").write_bytes(f"image-{index}".encode())
            client = FakeReadClient()
            active = 0
            max_active = 0
            uploaded: list[str] = []
            lock = threading.Lock()

            def fake_upload(_client, identity, data, timeout=60):
                nonlocal active, max_active
                with lock:
                    active += 1
                    max_active = max(max_active, active)
                    uploaded.append(identity["sha256"])
                time.sleep(0.03)
                with lock:
                    active -= 1
                return {
                    **identity,
                    "file_id": identity["sha256"][:12],
                    "url": f"https://sc04.alicdn.com/kf/H{identity['sha256'][:32]}.jpg",
                    "request_id": identity["sha256"][:8],
                    "success": True,
                    "raw_response": {},
                }

            workflow = Workflow(
                path,
                root / "artifacts",
                client=client,
                photo_cache_path=root / "photo-cache.json",
            )
            with patch("excel_workflow.upload_payload", side_effect=fake_upload):
                summary = workflow.prepare(upload_images=True)
            self.assertEqual(summary["unique_images"], 6)
            self.assertEqual(len(uploaded), 6)
            self.assertEqual(len(set(uploaded)), 6)
            self.assertEqual(max_active, 5)
            self.assertEqual(summary["ready"], 2)

    def test_workbook_uses_one_stable_backup_and_batch_save(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 2)
            ledger = WorkbookLedger(path)
            ledger.update("产品发布清单", 2, "待审核", "1701")
            ledger.update("产品发布清单", 3, "发布失败", summary="test")
            ledger.save()
            backups = list((root / ".publisher-backups").glob("*workflow-original.xlsx"))
            self.assertEqual(len(backups), 1)
            sheet = load_workbook(path)["产品发布清单"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers["发布状态"]).value, "待审核")
            self.assertEqual(sheet.cell(3, headers["发布状态"]).value, "发布失败")

    def seed_ready_items(self, workflow: Workflow, count: int,
                         ambiguous_index: int | None = None) -> None:
        workflow.state["catalog"] = {
            "complete": True,
            "fetched_at_epoch_ms": int(time.time() * 1000),
            "products": [],
        }
        for index in range(count):
            key = f"P{index + 1:03d}"
            title = "Ambiguous Product" if ambiguous_index == index else f"Add Product {index + 1}"
            product = ready_product(key, title)
            xml, report = build(SOURCE_XML, product)
            self.assertTrue(report["ready_for_submission"], report["errors"])
            item_dir = workflow.item_dir(key)
            (item_dir / "product.xml").write_bytes(xml.encode("utf-8"))
            workflow.state["items"][key] = {
                "task_key": key,
                "sheet": "产品发布清单",
                "row": index + 2,
                "product": product,
                "preflight": report,
                "status": "ready",
            }
        workflow.save_state()

    def test_schema_add_concurrency_is_two(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 4)
            client = FakeAddClient()
            workflow = Workflow(path, root / "artifacts", client=client)
            self.seed_ready_items(workflow, 4)
            summary = workflow.submit()
            self.assertEqual(summary["submitted"], 4)
            self.assertEqual(client.max_active, ADD_CONCURRENCY)
            self.assertFalse(summary["circuit_breaker"]["open"])
            self.assertEqual(summary["inventory_verified"], 4)
            self.assertEqual(client.inventory_update_calls, 4)
            sheet = load_workbook(path)["产品发布清单"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            for row in range(2, 6):
                self.assertEqual(sheet.cell(row, headers["发布状态"]).value, "已上传")
                self.assertTrue(sheet.cell(row, headers["新商品ID"]).value)

    def test_ambiguous_add_opens_breaker_and_stops_next_wave(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 4)
            client = FakeAddClient(ambiguous_title="Ambiguous Product")
            workflow = Workflow(path, root / "artifacts", client=client)
            self.seed_ready_items(workflow, 4, ambiguous_index=0)
            summary = workflow.submit()
            self.assertTrue(summary["circuit_breaker"]["open"])
            self.assertEqual(summary["remaining"], 2)
            self.assertEqual(sum(bool(item.get("attempt")) for item in workflow.state["items"].values()), 2)

    def test_uniform_inventory_is_initialized_even_when_render_is_temporarily_unavailable(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 1)
            client = FakeAddClient(render_available=False)
            workflow = Workflow(path, root / "artifacts", client=client)
            self.seed_ready_items(workflow, 1)
            summary = workflow.submit()
            self.assertEqual(summary["inventory_verified"], 1)
            self.assertEqual(client.inventory_update_calls, 1)
            sheet = load_workbook(path)["产品发布清单"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers["发布状态"]).value, "已上传")

    def test_pending_inventory_can_be_reconciled_without_repeating_add(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 1)
            client = FakeAddClient(inventory_available=False)
            workflow = Workflow(path, root / "artifacts", client=client)
            self.seed_ready_items(workflow, 1)
            submitted = workflow.submit()
            self.assertEqual(submitted["submitted"], 1)
            self.assertEqual(submitted["inventory_pending"], 1)
            self.assertEqual(client.counter, 1)
            client.inventory_available = True
            reconciled = workflow.reconcile_inventory()
            self.assertEqual(reconciled["verified"], 1)
            self.assertEqual(client.counter, 1)
            sheet = load_workbook(path)["产品发布清单"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers["发布状态"]).value, "已上传")

    def test_ambiguous_inventory_write_opens_breaker_and_is_not_retried(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 1)
            client = FakeAddClient(inventory_timeout=True)
            workflow = Workflow(path, root / "artifacts", client=client)
            self.seed_ready_items(workflow, 1)
            summary = workflow.submit()
            self.assertTrue(summary["circuit_breaker"]["open"])
            self.assertEqual(summary["inventory_unverified"], 1)
            self.assertEqual(client.inventory_update_calls, 1)
            with self.assertRaises(RuntimeError):
                workflow.reconcile_inventory()
            self.assertEqual(client.inventory_update_calls, 1)

    def test_watch_requires_exact_inventory_before_success(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = create_workbook(root, 1)
            product = ready_product("P001", "Verified Product")
            product["inventory_mode"] = "embedded"
            xml, report = build(SOURCE_XML, product)
            self.assertTrue(report["ready_for_submission"], report["errors"])
            rendered = ET.fromstring(xml)
            sku_id = rendered.find(".//field[@id='skuId']")
            value = ET.SubElement(sku_id, "value")
            value.text = "900001"
            render_xml = ET.tostring(rendered, encoding="unicode")
            client = FakeVerifyClient("1700000000001", product["title"], render_xml, 10)
            workflow = Workflow(path, root / "artifacts", client=client)
            item_dir = workflow.item_dir("P001")
            (item_dir / "product.xml").write_bytes(xml.encode("utf-8"))
            workflow.state["items"]["P001"] = {
                "task_key": "P001",
                "sheet": "产品发布清单",
                "row": 2,
                "product": product,
                "preflight": report,
                "status": "auditing",
                "attempt": {
                    "outcome": "accepted",
                    "product_id": "1700000000001",
                },
            }
            workflow.save_state()
            summary = workflow.watch(max_wait_seconds=1)
            self.assertEqual(summary["verified"], 1)
            sheet = load_workbook(path)["产品发布清单"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers["发布状态"]).value, "发布成功")


def main() -> None:
    global SOURCE_XML
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-render", required=True, type=Path)
    args = parser.parse_args()
    SOURCE_XML = extract_render_xml(args.source_render)
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(ExcelWorkflowTests)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    raise SystemExit(0 if result.wasSuccessful() else 1)


if __name__ == "__main__":
    main()
