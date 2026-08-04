#!/usr/bin/env python3
"""High-throughput, resumable Excel-driven Alibaba product publishing workflow."""

from __future__ import annotations

import argparse
import csv
from concurrent.futures import ThreadPoolExecutor, as_completed
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import threading
import time
import urllib.error
from typing import Any, Callable

from openpyxl import load_workbook

from api_client import AlibabaClient, load_environment, response_payload
from batch_tool import (
    catalog_title_matches,
    full_product_catalog,
    inventory_rows,
    inventory_targets_from_xml,
    response_product_id,
    selected_comparison,
    sku_ids_by_outer_id,
)
from clone_engine import build, extract_render_xml
from photobank import local_payload, upload_payload


READ_CONCURRENCY = 8
IMAGE_CONCURRENCY = 5
ADD_CONCURRENCY = 2
CATALOG_TTL_SECONDS = 15 * 60
CHECKPOINT_EVERY = 10
SCRIPT_VERSION = "2.1.0"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
INPUT_ALIASES = {
    "task_id": ("任务ID", "task_id"),
    "source_product_id": ("源商品ID", "source_product_id"),
    "title": ("新产品标题", "title"),
    "image_folder": ("新产品图片路径", "image_paths"),
    "specifications": ("新产品规格参数", "specifications"),
    "sku_plan_path": ("SKU方案文件路径", "sku_plan_path"),
    "inventory_mode": ("库存模式", "inventory_mode"),
    "allow_similar": ("允许相似商品", "allow_similar_listing"),
    "notes": ("备注", "notes"),
}
OUTPUT_COLUMNS = ("发布状态", "新商品ID", "验证时间", "验证摘要")
PENDING_STATUSES = {"auditing", "pending", "reviewing", "draft", ""}


def now_ms() -> int:
    return int(time.time() * 1000)


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def natural_key(path: Path) -> list[Any]:
    return [int(part) if part.isdigit() else part.casefold()
            for part in re.split(r"(\d+)", path.name)]


def atomic_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def read_json(path: Path, default: Any = None) -> Any:
    if not path.is_file():
        return default
    return json.loads(path.read_text(encoding="utf-8-sig"))


def normalize_title(value: str) -> str:
    return value.strip().casefold()


def resolve_path(raw: str, workbook: Path) -> Path:
    path = Path(raw.strip().strip('"'))
    return path if path.is_absolute() else workbook.parent / path


def resolve_image_folder(raw: str, workbook: Path) -> tuple[list[Path], list[str]]:
    if not raw:
        return [], ["新产品图片路径为空"]
    folder = resolve_path(raw, workbook)
    if not folder.is_dir():
        return [], [f"新产品图片路径必须是存在的文件夹: {folder}"]
    images = sorted(
        (path.resolve() for path in folder.iterdir()
         if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS),
        key=natural_key,
    )
    return images, [] if images else [f"图片文件夹内没有支持的图片: {folder}"]


def parse_pairs(raw: str) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for part in re.split(r"[;\r\n；]+", raw or ""):
        part = part.strip()
        if not part:
            continue
        if "=" not in part:
            raise ValueError(f"规格参数必须使用 名称=值: {part}")
        key, value = (piece.strip() for piece in part.split("=", 1))
        result[key] = [item.strip() for item in value.split("+")] if "+" in value else value
    return result


def parse_quantity_pairs(value: Any, value_key: str) -> list[dict]:
    if isinstance(value, list):
        return value
    result = []
    for part in re.split(r"[;,，；]+", text(value)):
        if not part:
            continue
        left, right = (item.strip() for item in part.split("=", 1))
        result.append({"quantity": int(left), value_key: right})
    return result


def parse_plan_text(path: Path) -> dict[str, Any]:
    raw = path.read_text(encoding="utf-8-sig")
    fenced = re.search(r"```json\s*(\{.*?\})\s*```", raw, re.I | re.S)
    if fenced:
        return json.loads(fenced.group(1))
    result: dict[str, Any] = {}
    for line in raw.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith(("#", "|", "```")) or ":" not in stripped:
            continue
        key, value = (part.strip() for part in stripped.split(":", 1))
        result[key] = value
    return result


def load_plan(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    elif suffix in {".md", ".markdown", ".txt"}:
        data = parse_plan_text(path)
    elif suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        data = rows[0] if rows else {}
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.active
        headers = [text(cell.value) for cell in sheet[1]]
        values = [text(cell.value) for cell in sheet[2]] if sheet.max_row >= 2 else []
        data = dict(zip(headers, values))
    else:
        raise ValueError(f"不支持的SKU方案格式: {path.suffix}")
    if not isinstance(data, dict):
        raise ValueError("SKU方案必须解析为对象。")
    return normalize_plan(data)


def normalize_plan(data: dict[str, Any]) -> dict[str, Any]:
    plan = dict(data)
    aliases = {
        "variants": ("variants", "规格值", "颜色", "colors"),
        "variant_axis": ("variant_axis", "规格轴", "销售属性"),
        "price_tiers": ("price_tiers", "阶梯价格"),
        "moq": ("moq", "MOQ", "起订量"),
        "lead_time_days": ("lead_time_days", "交期", "交期天数"),
        "stock": ("stock", "库存"),
        "warehouse_code": ("warehouse_code", "仓库代码"),
        "pkg_length": ("pkg_length", "包装长"),
        "pkg_width": ("pkg_width", "包装宽"),
        "pkg_height": ("pkg_height", "包装高"),
        "pkg_weight": ("pkg_weight", "包装重量"),
    }
    normalized: dict[str, Any] = {}
    for target, names in aliases.items():
        value = next((plan[name] for name in names if name in plan and text(plan[name])), None)
        if value is not None:
            normalized[target] = value
    for key in (
        "sale_properties", "skus", "category_attributes", "category_attribute_overrides",
        "custom_properties", "price_tiers", "lead_times", "package", "sku_defaults",
        "variant_axes", "simple_fields", "omit_category_attributes", "omit_top_level_fields",
        "inventory_mode", "super_text",
    ):
        if key in plan:
            normalized[key] = plan[key]
    if "price_tiers" in normalized:
        normalized["price_tiers"] = parse_quantity_pairs(normalized["price_tiers"], "price")
    if "lead_times" in normalized:
        normalized["lead_times"] = parse_quantity_pairs(normalized["lead_times"], "day")
    if "variant_axes" not in normalized and text(normalized.get("variants")):
        variants = [part.strip() for part in re.split(r"[;,，；]+", text(normalized["variants"])) if part.strip()]
        normalized["variant_axes"] = [{
            "name": text(normalized.get("variant_axis")) or "Color",
            "values": variants,
        }]
    if "sku_defaults" not in normalized and ("stock" in normalized or "warehouse_code" in normalized):
        normalized["sku_defaults"] = {
            "stock_target": int(normalized.get("stock") or 0),
            "warehouse_code": text(normalized.get("warehouse_code")) or "CN_LOCAL_01",
        }
    if "moq" in normalized:
        normalized["moq"] = int(normalized["moq"])
    if "lead_times" not in normalized and "lead_time_days" in normalized and "moq" in normalized:
        normalized["lead_times"] = [{
            "quantity": int(normalized["moq"]),
            "day": int(normalized["lead_time_days"]),
        }]
    if "package" not in normalized and all(key in normalized for key in ("pkg_length", "pkg_width", "pkg_height", "pkg_weight")):
        normalized["package"] = {
            "length": float(normalized["pkg_length"]),
            "width": float(normalized["pkg_width"]),
            "height": float(normalized["pkg_height"]),
            "weight": float(normalized["pkg_weight"]),
        }
    return normalized


class WorkbookLedger:
    def __init__(self, path: Path, sheet_name: str | None = None):
        self.path = path.resolve()
        self.keep_vba = self.path.suffix.lower() == ".xlsm"
        self.workbook = load_workbook(self.path, data_only=False, keep_vba=self.keep_vba)
        self.sheets = [self.workbook[sheet_name]] if sheet_name else self.workbook.worksheets
        self.output_columns: dict[str, dict[str, int]] = {}
        self.dirty = 0
        self.backup_path = self._ensure_backup()
        for sheet in self.sheets:
            headers = {text(cell.value): cell.column for cell in sheet[1] if text(cell.value)}
            for name in OUTPUT_COLUMNS:
                if name not in headers:
                    column = sheet.max_column + 1
                    sheet.cell(1, column, name)
                    headers[name] = column
            self.output_columns[sheet.title] = headers

    def _ensure_backup(self) -> Path:
        backup_dir = self.path.parent / ".publisher-backups"
        backup_dir.mkdir(exist_ok=True)
        marker = backup_dir / f"{self.path.stem}-workflow-original{self.path.suffix}"
        if not marker.exists():
            shutil.copy2(self.path, marker)
        return marker

    def update(self, sheet_name: str, row: int, status: str, product_id: str = "",
               summary: str = "", verified_at: str = "") -> None:
        sheet = self.workbook[sheet_name]
        columns = self.output_columns[sheet_name]
        sheet.cell(row, columns["发布状态"], status)
        if product_id:
            existing = text(sheet.cell(row, columns["新商品ID"]).value)
            if existing and existing != product_id:
                raise RuntimeError(f"{sheet_name}!{row} already contains another product ID.")
            sheet.cell(row, columns["新商品ID"], product_id)
        if verified_at:
            sheet.cell(row, columns["验证时间"], verified_at)
        if summary:
            sheet.cell(row, columns["验证摘要"], summary[:500])
        self.dirty += 1
        if self.dirty >= CHECKPOINT_EVERY:
            self.save()

    def save(self) -> None:
        if not self.dirty:
            return
        temporary = self.path.with_name(f".{self.path.stem}.workflow{self.path.suffix}")
        self.workbook.save(temporary)
        temporary.replace(self.path)
        self.dirty = 0


def header_columns(sheet) -> tuple[dict[str, int], list[str]]:
    headers = {text(cell.value): cell.column for cell in sheet[1] if text(cell.value)}
    columns: dict[str, int] = {}
    missing = []
    for field, aliases in INPUT_ALIASES.items():
        match = next((headers[name] for name in aliases if name in headers), None)
        if match:
            columns[field] = match
        elif field in {"source_product_id", "title", "image_folder", "specifications", "sku_plan_path"}:
            missing.append(aliases[0])
    return columns, missing


def compile_excel_rows(input_path: Path, sheet_name: str | None = None) -> list[dict]:
    keep_vba = input_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(input_path, data_only=False, keep_vba=keep_vba)
    sheets = [workbook[sheet_name]] if sheet_name else workbook.worksheets
    rows: list[dict] = []
    for sheet in sheets:
        columns, missing = header_columns(sheet)
        if missing:
            if sheet_name:
                raise ValueError(f"{sheet.title} 缺少列: {', '.join(missing)}")
            continue
        for row_number in range(2, sheet.max_row + 1):
            values = {name: text(sheet.cell(row_number, column).value)
                      for name, column in columns.items()}
            if not any(values.values()):
                continue
            task_key = values.get("task_id") or f"{sheet.title}!{row_number}"
            errors = []
            source_id = values.get("source_product_id", "")
            if not source_id.isdigit():
                errors.append("源商品ID必须是数字")
            if not values.get("title"):
                errors.append("新产品标题为空")
            images, image_errors = resolve_image_folder(values.get("image_folder", ""), input_path)
            errors.extend(image_errors)
            plan_path = resolve_path(values.get("sku_plan_path", ""), input_path)
            plan = {}
            if not plan_path.is_file():
                errors.append(f"SKU方案文件不存在: {plan_path}")
            else:
                try:
                    plan = load_plan(plan_path)
                except Exception as exc:
                    errors.append(f"SKU方案解析失败: {exc}")
            try:
                specs = parse_pairs(values.get("specifications", ""))
            except ValueError as exc:
                specs = {}
                errors.append(str(exc))
            product = {
                **plan,
                "item_key": task_key,
                "source_product_id": source_id,
                "title": values.get("title", ""),
                "inventory_mode": values.get("inventory_mode") or plan.get("inventory_mode") or "embedded",
                "category_attribute_overrides": {
                    **(plan.get("category_attribute_overrides") or {}),
                    **specs,
                },
                "_image_paths": [str(path) for path in images],
            }
            rows.append({
                "task_key": task_key,
                "sheet": sheet.title,
                "row": row_number,
                "product": product,
                "errors": errors,
            })
    if not rows:
        raise ValueError("Excel中没有商品数据行。")
    titles = [normalize_title(item["product"]["title"]) for item in rows if item["product"]["title"]]
    if len(titles) != len(set(titles)):
        duplicate_titles = sorted({title for title in titles if titles.count(title) > 1})
        for item in rows:
            if normalize_title(item["product"]["title"]) in duplicate_titles:
                item["errors"].append("Excel内部存在重复标题")
    return rows


class Workflow:
    def __init__(self, input_path: Path, work_dir: Path, sheet: str | None = None,
                 mcp_config: Path | None = None, client: Any | None = None,
                 photo_cache_path: Path | None = None):
        self.input_path = input_path.resolve()
        self.root = work_dir.resolve() / self.input_path.stem
        self.root.mkdir(parents=True, exist_ok=True)
        self.state_path = self.root / "workflow-state.json"
        self.state_lock = threading.Lock()
        self.client = client or AlibabaClient(load_environment(mcp_config))
        self.sheet = sheet
        self.photo_cache_path = photo_cache_path or (
            Path.home() / ".codex" / "cache" / "alibaba-batch-product-publisher" / "photobank-map.json"
        )
        self.state = read_json(self.state_path, {
            "version": SCRIPT_VERSION,
            "workbook": str(self.input_path),
            "workbook_sha256_at_start": hashlib.sha256(self.input_path.read_bytes()).hexdigest(),
            "created_at_epoch_ms": now_ms(),
            "catalog": {},
            "items": {},
            "circuit_breaker": {"open": False},
        })

    def save_state(self) -> None:
        with self.state_lock:
            atomic_json(self.state_path, self.state)

    def item_dir(self, task_key: str) -> Path:
        safe = re.sub(r'[<>:"/\\|?*]+', "_", task_key)
        path = self.root / "items" / safe
        path.mkdir(parents=True, exist_ok=True)
        return path

    def source_cache_dir(self, source_id: str) -> Path:
        path = self.root / "source-cache" / source_id
        path.mkdir(parents=True, exist_ok=True)
        return path

    def fetch_source(self, source_id: str) -> dict:
        cache_dir = self.source_cache_dir(source_id)
        product_response = self.client.request("alibaba.icbu.product.get", {
            "product_id": source_id,
            "language": "ENGLISH",
        })
        atomic_json(cache_dir / "product-get.json", product_response)
        source = response_payload(product_response).get("product") or {}
        category_id = source.get("category_id") or source.get("categoryId")
        if not category_id:
            raise RuntimeError(f"{source_id}: cannot resolve source category.")
        modified = text(source.get("gmt_modified") or source.get("gmtModified"))
        metadata = read_json(cache_dir / "metadata.json", {})
        render_file = cache_dir / "source-render.xml"
        render_response_file = cache_dir / "source-render-response.json"
        if metadata.get("source_modified") == modified and render_file.is_file():
            source_xml = render_file.read_text(encoding="utf-8-sig")
            reused = True
        else:
            response = self.client.request("alibaba.icbu.product.schema.render", {
                "param_product_top_publish_request": {
                    "product_id": source_id,
                    "cat_id": int(category_id),
                    "language": "en_US",
                }
            })
            atomic_json(render_response_file, response)
            payload = response_payload(response)
            if not payload.get("biz_success") or not payload.get("data"):
                raise RuntimeError(f"{source_id}: source Render failed.")
            source_xml = str(payload["data"])
            render_file.write_text(source_xml, encoding="utf-8")
            reused = False
        metadata = {
            "source_product_id": source_id,
            "category_id": int(category_id),
            "source_modified": modified,
            "xml_sha256": hashlib.sha256(source_xml.encode("utf-8")).hexdigest(),
            "render_reused": reused,
            "fetched_at_epoch_ms": now_ms(),
        }
        schema_response = self.client.request("alibaba.icbu.product.schema.get", {
            "param_product_top_publish_request": {
                "cat_id": int(category_id),
                "language": "en_US",
            }
        })
        atomic_json(cache_dir / "current-schema-response.json", schema_response)
        schema_payload = response_payload(schema_response)
        current_schema_xml = (
            schema_payload.get("data")
            or schema_payload.get("item_schema")
            or schema_payload.get("xml")
        )
        if not schema_payload.get("biz_success") or not current_schema_xml:
            raise RuntimeError(f"{source_id}: current category Schema failed.")
        current_schema_xml = str(current_schema_xml)
        (cache_dir / "current-schema.xml").write_text(current_schema_xml, encoding="utf-8")
        metadata["current_schema_sha256"] = hashlib.sha256(
            current_schema_xml.encode("utf-8")
        ).hexdigest()
        atomic_json(cache_dir / "metadata.json", metadata)
        return {
            "category_id": int(category_id),
            "source_xml": source_xml,
            "current_schema_xml": current_schema_xml,
            "metadata": metadata,
        }

    def fetch_catalog(self, evidence_name: str) -> list[dict]:
        evidence = self.root / evidence_name
        catalog = full_product_catalog(self.client, evidence)
        snapshot = {
            "complete": True,
            "fetched_at_epoch_ms": now_ms(),
            "products": catalog,
            "title_index": {
                normalize_title(item["title"]): item["product_id"] for item in catalog
            },
        }
        self.state["catalog"] = snapshot
        self.save_state()
        return catalog

    def catalog(self, refresh_if_stale: bool) -> list[dict]:
        snapshot = self.state.get("catalog") or {}
        age = (now_ms() - int(snapshot.get("fetched_at_epoch_ms") or 0)) / 1000
        if not snapshot.get("complete") or (refresh_if_stale and age > CATALOG_TTL_SECONDS):
            return self.fetch_catalog("duplicate-catalog-refresh" if snapshot else "duplicate-catalog")
        return list(snapshot["products"])

    def prepare(self, upload_images: bool) -> dict:
        if any((item.get("attempt") for item in self.state.get("items", {}).values())):
            raise RuntimeError("已有商品新增尝试记录；禁止重新prepare覆盖已锁定的XML。")
        rows = compile_excel_rows(self.input_path, self.sheet)
        ledger = WorkbookLedger(self.input_path, self.sheet)
        source_ids = sorted({row["product"]["source_product_id"] for row in rows if not row["errors"]})
        sources: dict[str, dict] = {}
        source_errors: dict[str, str] = {}
        catalog_future = None
        with ThreadPoolExecutor(max_workers=READ_CONCURRENCY) as executor:
            catalog_future = executor.submit(self.fetch_catalog, "duplicate-catalog")
            future_sources = {executor.submit(self.fetch_source, source_id): source_id for source_id in source_ids}
            for future in as_completed(future_sources):
                source_id = future_sources[future]
                try:
                    sources[source_id] = future.result()
                except Exception as exc:
                    source_errors[source_id] = f"{type(exc).__name__}: {exc}"
            catalog = catalog_future.result()
        title_index = {normalize_title(item["title"]): item for item in catalog}

        payloads: dict[str, tuple[dict, bytes]] = {}
        image_uses: dict[str, list[str]] = {}
        for row in rows:
            for raw_path in row["product"].pop("_image_paths", []):
                try:
                    identity, data = local_payload(Path(raw_path))
                    sha = identity["sha256"]
                    payloads.setdefault(sha, (identity, data))
                    image_uses.setdefault(sha, []).append(row["task_key"])
                    row["product"].setdefault("_image_hashes", []).append(sha)
                except Exception as exc:
                    row["errors"].append(f"图片读取失败: {exc}")

        photo_cache_path = self.photo_cache_path
        photo_cache = read_json(photo_cache_path, {"images": {}})
        upload_results: dict[str, dict] = {}
        to_upload: dict[str, tuple[dict, bytes]] = {}
        for sha, payload in payloads.items():
            cached = (photo_cache.get("images") or {}).get(sha) or {}
            if (
                cached.get("success")
                and cached.get("file_id")
                and "alicdn.com/" in text(cached.get("url"))
                and cached.get("request_id")
            ):
                upload_results[sha] = cached
            elif upload_images:
                if cached.get("attempted") and not cached.get("success"):
                    upload_results[sha] = {
                        **cached,
                        "success": False,
                        "error": "存在不完整上传记录，禁止自动重试",
                    }
                else:
                    to_upload[sha] = payload

        if to_upload:
            cache_lock = threading.Lock()

            def do_upload(sha: str, payload: tuple[dict, bytes]) -> tuple[str, dict]:
                identity, data = payload
                started = {
                    **identity,
                    "attempted": True,
                    "success": False,
                    "attempted_at_epoch_ms": now_ms(),
                }
                with cache_lock:
                    photo_cache.setdefault("images", {})[sha] = started
                    atomic_json(photo_cache_path, photo_cache)
                try:
                    result = upload_payload(self.client, identity, data)
                    complete = {**started, **result, "attempted": True}
                except Exception as exc:
                    complete = {
                        **started,
                        "error": f"{type(exc).__name__}: {exc}",
                        "ambiguous": True,
                    }
                with cache_lock:
                    photo_cache["images"][sha] = complete
                    atomic_json(photo_cache_path, photo_cache)
                return sha, complete

            with ThreadPoolExecutor(max_workers=IMAGE_CONCURRENCY) as executor:
                futures = {executor.submit(do_upload, sha, payload): sha for sha, payload in to_upload.items()}
                for future in as_completed(futures):
                    sha, result = future.result()
                    upload_results[sha] = result

        prepared_rows: list[dict] = []

        def build_row(row: dict) -> dict:
            item_dir = self.item_dir(row["task_key"])
            product = row["product"]
            errors = list(row["errors"])
            source_id = product["source_product_id"]
            if source_id in source_errors:
                errors.append(f"源商品读取失败: {source_errors[source_id]}")
            source = sources.get(source_id)
            if source:
                product["category_id"] = source["category_id"]
            hashes = product.pop("_image_hashes", [])
            resolved_images = []
            for sha in hashes:
                image = upload_results.get(sha)
                if not image or not image.get("success"):
                    errors.append(f"图片未解析: {sha} {(image or {}).get('error', '')}")
                    continue
                resolved_images.append({"file_id": image["file_id"], "url": image["url"], "local_sha256": sha})
            if not upload_images and hashes:
                errors.append("图片尚未上传；重新运行 prepare --upload-images。")
            if resolved_images:
                product["main_images"] = resolved_images[:6]
                product["detail_galleries"] = [{
                    "gallery": "200",
                    "display_name": "Scene image",
                    "images": [{"url": image["url"]} for image in resolved_images],
                }]
            duplicate = title_index.get(normalize_title(product["title"]))
            if duplicate:
                errors.append(f"精确标题已存在: {duplicate['product_id']}")
            xml = ""
            report = {"errors": errors, "ready_for_submission": False}
            if source:
                xml, report = build(
                    source["source_xml"],
                    product,
                    current_schema_xml=source["current_schema_xml"],
                )
                report["errors"] = errors + list(report.get("errors") or [])
                report["duplicate_check"] = {
                    "performed": True,
                    "exact_matches": [duplicate] if duplicate else [],
                }
                report["ready_for_submission"] = not report["errors"]
                xml_path = item_dir / "product.xml"
                xml_path.write_bytes(xml.encode("utf-8"))
                report["xml_file"] = str(xml_path)
                report["source_render_file"] = str(self.source_cache_dir(source_id) / "source-render.xml")
                report["current_schema_file"] = str(self.source_cache_dir(source_id) / "current-schema.xml")
            atomic_json(item_dir / "preflight.json", report)
            atomic_json(item_dir / "product.json", product)
            return {
                "task_key": row["task_key"],
                "sheet": row["sheet"],
                "row": row["row"],
                "product": product,
                "preflight": report,
                "status": "ready" if report.get("ready_for_submission") else "blocked",
            }

        with ThreadPoolExecutor(max_workers=READ_CONCURRENCY) as executor:
            futures = [executor.submit(build_row, row) for row in rows]
            for future in as_completed(futures):
                prepared_rows.append(future.result())
        prepared_rows.sort(key=lambda item: (item["sheet"], item["row"]))
        for item in prepared_rows:
            self.state["items"][item["task_key"]] = item
            if item["status"] == "ready":
                ledger.update(item["sheet"], item["row"], "预检完成", summary="图片、源结构、标题与XML预检通过")
            else:
                summary = "; ".join(item["preflight"].get("errors") or [])[:500]
                ledger.update(item["sheet"], item["row"], "预检失败", summary=summary)
        ledger.save()
        self.state["prepared_at_epoch_ms"] = now_ms()
        self.save_state()
        summary = {
            "ready": sum(item["status"] == "ready" for item in prepared_rows),
            "blocked": sum(item["status"] != "ready" for item in prepared_rows),
            "unique_sources": len(source_ids),
            "unique_images": len(payloads),
            "uploaded_or_reused_images": sum(bool(result.get("success")) for result in upload_results.values()),
            "state": str(self.state_path),
        }
        atomic_json(self.root / "prepare-summary.json", summary)
        return summary

    def submit_one(self, key: str, catalog: list[dict]) -> dict:
        item = self.state["items"][key]
        product = item["product"]
        report = item["preflight"]
        item_dir = self.item_dir(key)
        if item.get("attempt"):
            return {"task_key": key, "outcome": "skipped_existing_attempt"}
        matches = catalog_title_matches(catalog, product["title"])
        if matches:
            item["status"] = "blocked"
            item["preflight"]["errors"].append("正式提交前发现精确同标题商品。")
            self.save_state()
            return {"task_key": key, "outcome": "duplicate_blocked"}
        xml_path = item_dir / "product.xml"
        actual_hash = hashlib.sha256(xml_path.read_bytes()).hexdigest()
        if actual_hash != report.get("xml_sha256"):
            item["status"] = "blocked"
            item["preflight"]["errors"].append("XML哈希与预检不一致。")
            self.save_state()
            return {"task_key": key, "outcome": "hash_blocked"}
        attempt = {
            "api": "alibaba.icbu.product.schema.add",
            "title": product["title"],
            "category_id": product["category_id"],
            "xml_sha256": actual_hash,
            "attempted_at_epoch_ms": now_ms(),
            "outcome": "attempt_started",
            "automatic_retry_allowed": False,
        }
        with self.state_lock:
            if self.state["items"][key].get("attempt"):
                return {"task_key": key, "outcome": "skipped_existing_attempt"}
            self.state["items"][key]["attempt"] = attempt
            atomic_json(self.state_path, self.state)
        atomic_json(item_dir / "add-attempt.json", attempt)
        try:
            response = self.client.request("alibaba.icbu.product.schema.add", {
                "param_product_top_publish_request": {
                    "cat_id": int(product["category_id"]),
                    "language": "en_US",
                    "xml": xml_path.read_text(encoding="utf-8"),
                }
            })
            atomic_json(item_dir / "add-response.json", response)
            payload = response_payload(response)
            product_id = response_product_id(response)
            accepted = bool(payload.get("biz_success") and product_id)
            attempt.update({
                "outcome": "accepted" if accepted else "api_rejected",
                "product_id": product_id,
                "request_id": payload.get("request_id"),
                "trace_id": payload.get("trace_id") or payload.get("_trace_id_"),
                "msg_code": payload.get("msg_code"),
            })
            item["status"] = "auditing" if accepted else "rejected"
        except Exception as exc:
            attempt.update({
                "outcome": "ambiguous_exception",
                "exception_type": type(exc).__name__,
            })
            item["status"] = "ambiguous"
            self.state["circuit_breaker"] = {
                "open": True,
                "task_key": key,
                "opened_at_epoch_ms": now_ms(),
                "reason": type(exc).__name__,
            }
            atomic_json(item_dir / "add-exception.json", attempt)
        atomic_json(item_dir / "add-attempt.json", attempt)
        self.save_state()
        return {"task_key": key, **attempt}

    def sync_initial_inventory(self, key: str) -> dict:
        """Set and verify real SKU inventory after add; never repeat an uncertain delta."""
        item = self.state["items"][key]
        product = item["product"]
        product_id = text((item.get("attempt") or {}).get("product_id"))
        targets = item.get("preflight", {}).get("inventory_targets") or []
        item_dir = self.item_dir(key)
        if not product_id or not targets:
            item["inventory_state"] = "not_required"
            self.save_state()
            return {"task_key": key, "outcome": "not_required", "verified": True}
        if self.state.get("circuit_breaker", {}).get("open"):
            item["inventory_state"] = "pending"
            self.save_state()
            return {"task_key": key, "outcome": "pending", "verified": False}
        if item.get("inventory_attempt"):
            prior = item["inventory_attempt"]
            return {
                "task_key": key,
                "outcome": prior.get("outcome") or "existing_attempt",
                "verified": prior.get("outcome") == "verified",
            }

        try:
            render_response = self.client.request("alibaba.icbu.product.schema.render", {
                "param_product_top_publish_request": {
                    "product_id": product_id,
                    "cat_id": int(product["category_id"]),
                    "language": "en_US",
                }
            })
            inventory_response = self.client.request("alibaba.icbu.product.sku.inventory.get", {
                "language": "en_US", "product_id": product_id,
            })
        except Exception as exc:
            item["inventory_state"] = "pending"
            atomic_json(item_dir / "initial-inventory-read-error.json", {
                "error_type": type(exc).__name__,
                "write_attempted": False,
            })
            self.save_state()
            return {"task_key": key, "outcome": "pending", "verified": False}

        atomic_json(item_dir / "initial-inventory-render.json", render_response)
        atomic_json(item_dir / "initial-inventory-before.json", inventory_response)
        render_xml = text(response_payload(render_response).get("data"))
        rows = inventory_rows(inventory_response)
        sku_map = sku_ids_by_outer_id(render_xml) if render_xml else {}
        expected: dict[tuple[str, str], int] = {}
        for target in targets:
            sku_id = sku_map.get(text(target.get("sku_outer_id")))
            warehouse = text(target.get("warehouse_code"))
            if sku_id and warehouse and target.get("target") is not None:
                expected[(sku_id, warehouse)] = int(target["target"])

        # Immediately after add, Render may be unavailable during review. A uniform target
        # is still safe to map directly to every returned SKU row because no SKU-specific
        # distinction is required.
        target_values = {int(target["target"]) for target in targets if target.get("target") is not None}
        warehouses = {text(target.get("warehouse_code")) for target in targets if target.get("warehouse_code")}
        if len(expected) != len(targets) and len(rows) == len(targets) and len(target_values) == 1 and len(warehouses) == 1:
            target_value = next(iter(target_values))
            warehouse = next(iter(warehouses))
            if all(text(row.get("sku_id")) and text(row.get("inventory_code")) == warehouse for row in rows):
                expected = {(text(row["sku_id"]), warehouse): target_value for row in rows}

        actual = {
            (text(row.get("sku_id")), text(row.get("inventory_code"))): int(row.get("inventory") or 0)
            for row in rows
        }
        if len(expected) != len(targets) or any(pair not in actual for pair in expected):
            item["inventory_state"] = "pending"
            self.save_state()
            return {"task_key": key, "outcome": "pending", "verified": False}

        changes = [{
            "sku_id": int(sku_id),
            "inventory_code": warehouse,
            "inventory": abs(target - actual[(sku_id, warehouse)]),
            "operate": "plus" if target > actual[(sku_id, warehouse)] else "sub",
        } for (sku_id, warehouse), target in expected.items() if actual[(sku_id, warehouse)] != target]
        attempt = {
            "api": "alibaba.icbu.product.inventory.update",
            "product_id": product_id,
            "before": {f"{sku_id}:{warehouse}": actual[(sku_id, warehouse)] for sku_id, warehouse in expected},
            "targets": {f"{sku_id}:{warehouse}": target for (sku_id, warehouse), target in expected.items()},
            "changes": changes,
            "attempted_at_epoch_ms": now_ms(),
            "outcome": "attempt_started" if changes else "verified",
            "automatic_retry_allowed": False,
        }
        item["inventory_attempt"] = attempt
        atomic_json(item_dir / "initial-inventory-attempt.json", attempt)
        self.save_state()
        if not changes:
            item["inventory_state"] = "verified"
            self.save_state()
            return {"task_key": key, "outcome": "verified", "verified": True}

        try:
            update_response = self.client.request("alibaba.icbu.product.inventory.update", {
                "request_param": {
                    "product_id": int(product_id),
                    "inventory_list": changes,
                }
            })
            atomic_json(item_dir / "initial-inventory-update-response.json", update_response)
        except Exception as exc:
            attempt.update({"outcome": "ambiguous_exception", "exception_type": type(exc).__name__})
            item["inventory_state"] = "ambiguous"
            self.state["circuit_breaker"] = {
                "open": True,
                "task_key": key,
                "opened_at_epoch_ms": now_ms(),
                "reason": f"inventory:{type(exc).__name__}",
            }
            atomic_json(item_dir / "initial-inventory-attempt.json", attempt)
            self.save_state()
            return {"task_key": key, "outcome": "ambiguous_exception", "verified": False}

        after: dict[tuple[str, str], int] = {}
        try:
            for poll in range(5):
                if poll:
                    time.sleep(1)
                latest = self.client.request("alibaba.icbu.product.sku.inventory.get", {
                    "language": "en_US", "product_id": product_id,
                })
                atomic_json(item_dir / f"initial-inventory-after-{poll + 1}.json", latest)
                after = {
                    (text(row.get("sku_id")), text(row.get("inventory_code"))): int(row.get("inventory") or 0)
                    for row in inventory_rows(latest)
                }
                if all(after.get(pair) == target for pair, target in expected.items()):
                    break
        except Exception as exc:
            attempt["readback_error_type"] = type(exc).__name__
        verified = bool(after) and all(after.get(pair) == target for pair, target in expected.items())
        attempt.update({
            "after": {f"{sku_id}:{warehouse}": after.get((sku_id, warehouse)) for sku_id, warehouse in expected},
            "outcome": "verified" if verified else "unverified",
        })
        item["inventory_state"] = "verified" if verified else "unverified"
        if not verified:
            self.state["circuit_breaker"] = {
                "open": True,
                "task_key": key,
                "opened_at_epoch_ms": now_ms(),
                "reason": "inventory_readback_unverified",
            }
        atomic_json(item_dir / "initial-inventory-attempt.json", attempt)
        self.save_state()
        return {"task_key": key, "outcome": attempt["outcome"], "verified": verified}

    def submit(self) -> dict:
        if self.state.get("circuit_breaker", {}).get("open"):
            raise RuntimeError("写入熔断器已打开；必须先只读核对不明确结果。")
        catalog = self.catalog(refresh_if_stale=True)
        eligible = [
            key for key, item in self.state["items"].items()
            if item.get("status") == "ready" and not item.get("attempt")
        ]
        ledger = WorkbookLedger(self.input_path, self.sheet)
        results = []
        for start in range(0, len(eligible), ADD_CONCURRENCY):
            wave = eligible[start:start + ADD_CONCURRENCY]
            with ThreadPoolExecutor(max_workers=ADD_CONCURRENCY) as executor:
                futures = [executor.submit(self.submit_one, key, catalog) for key in wave]
                wave_results = [future.result() for future in as_completed(futures)]
            results.extend(wave_results)
            for result in wave_results:
                item = self.state["items"][result["task_key"]]
                product_id = text((item.get("attempt") or {}).get("product_id"))
                if item["status"] == "auditing":
                    inventory = self.sync_initial_inventory(result["task_key"])
                    if inventory["verified"]:
                        ledger.update(
                            item["sheet"], item["row"], "已上传", product_id,
                            "新增接口已接受；真实SKU库存已同步并核验；未等待平台审核",
                            time.strftime("%Y-%m-%dT%H:%M:%S%z"),
                        )
                    elif inventory["outcome"] == "pending":
                        item["status"] = "inventory_pending"
                        ledger.update(
                            item["sheet"], item["row"], "已上传（库存待同步）", product_id,
                            "新增接口已接受；尚未取得可安全映射的真实SKU库存，禁止视为发品完成",
                            time.strftime("%Y-%m-%dT%H:%M:%S%z"),
                        )
                    else:
                        item["status"] = "inventory_unverified"
                        ledger.update(
                            item["sheet"], item["row"], "库存结果不明确", product_id,
                            "库存写入或回读未明确验证；已熔断，禁止自动重试",
                            time.strftime("%Y-%m-%dT%H:%M:%S%z"),
                        )
                    catalog.append({"product_id": product_id, "title": item["product"]["title"], "status": "auditing"})
                elif item["status"] == "ambiguous":
                    ledger.update(item["sheet"], item["row"], "提交结果不明确", summary="写入熔断已启动，禁止自动重试")
                elif item["status"] == "rejected":
                    ledger.update(item["sheet"], item["row"], "发布失败", summary="新增接口明确拒绝")
                elif item["status"] == "blocked":
                    ledger.update(item["sheet"], item["row"], "预检失败", summary="提交前标题或XML哈希检查失败")
            self.save_state()
            ledger.save()
            if self.state.get("circuit_breaker", {}).get("open"):
                break
        summary = {
            "submitted": sum(result.get("outcome") == "accepted" for result in results),
            "rejected": sum(result.get("outcome") == "api_rejected" for result in results),
            "ambiguous": sum(result.get("outcome") == "ambiguous_exception" for result in results),
            "inventory_verified": sum(item.get("inventory_state") == "verified" for item in self.state["items"].values()),
            "inventory_pending": sum(item.get("inventory_state") == "pending" for item in self.state["items"].values()),
            "inventory_unverified": sum(item.get("inventory_state") in {"ambiguous", "unverified"} for item in self.state["items"].values()),
            "remaining": sum(
                item.get("status") == "ready" and not item.get("attempt")
                for item in self.state["items"].values()
            ),
            "circuit_breaker": self.state.get("circuit_breaker"),
        }
        atomic_json(self.root / "submit-summary.json", summary)
        return summary

    def reconcile_inventory(self) -> dict:
        """Retry only pending read/mapping work; never retry a recorded inventory write."""
        if self.state.get("circuit_breaker", {}).get("open"):
            raise RuntimeError("写入熔断器已打开；禁止继续库存写入。")
        ledger = WorkbookLedger(self.input_path, self.sheet)
        keys = [
            key for key, item in self.state["items"].items()
            if item.get("inventory_state") == "pending"
            and (item.get("attempt") or {}).get("outcome") == "accepted"
            and not item.get("inventory_attempt")
        ]
        results = []
        for key in keys:
            result = self.sync_initial_inventory(key)
            results.append(result)
            item = self.state["items"][key]
            product_id = text((item.get("attempt") or {}).get("product_id"))
            if result["verified"]:
                item["status"] = "auditing"
                ledger.update(
                    item["sheet"], item["row"], "已上传", product_id,
                    "新增接口已接受；真实SKU库存已同步并核验；未等待平台审核",
                    time.strftime("%Y-%m-%dT%H:%M:%S%z"),
                )
            elif result["outcome"] != "pending":
                item["status"] = "inventory_unverified"
                ledger.update(
                    item["sheet"], item["row"], "库存结果不明确", product_id,
                    "库存写入或回读未明确验证；已熔断，禁止自动重试",
                    time.strftime("%Y-%m-%dT%H:%M:%S%z"),
                )
            if self.state.get("circuit_breaker", {}).get("open"):
                break
        ledger.save()
        self.save_state()
        summary = {
            "processed": len(results),
            "verified": sum(result.get("verified") is True for result in results),
            "pending": sum(result.get("outcome") == "pending" for result in results),
            "unverified": sum(result.get("outcome") not in {"pending", "verified"} for result in results),
            "circuit_breaker": self.state.get("circuit_breaker"),
        }
        atomic_json(self.root / "inventory-reconcile-summary.json", summary)
        return summary

    def verify_item(self, key: str) -> dict:
        item = self.state["items"][key]
        product = item["product"]
        product_id = text((item.get("attempt") or {}).get("product_id"))
        item_dir = self.item_dir(key)
        if not product_id:
            return {"task_key": key, "verified": False, "pending": False, "reason": "No product ID"}
        get_response = self.client.request("alibaba.icbu.product.get", {
            "product_id": product_id, "language": "ENGLISH",
        })
        render_response = self.client.request("alibaba.icbu.product.schema.render", {
            "param_product_top_publish_request": {
                "product_id": product_id,
                "cat_id": int(product["category_id"]),
                "language": "en_US",
            }
        })
        atomic_json(item_dir / "verify-get.json", get_response)
        atomic_json(item_dir / "verify-render.json", render_response)
        get_product = response_payload(get_response).get("product") or {}
        render_payload = response_payload(render_response)
        status = text(get_product.get("status")).lower()
        display = text(get_product.get("display")).upper()
        pending = status in PENDING_STATUSES or status != "approved" or display != "Y"
        expected_xml = (item_dir / "product.xml").read_text(encoding="utf-8")
        render_xml = text(render_payload.get("data"))
        fields_match = bool(render_payload.get("biz_success") and render_xml) and (
            selected_comparison(expected_xml) == selected_comparison(render_xml)
        )
        inventory_targets = item.get("preflight", {}).get("inventory_targets") or []
        inventory_matches = not inventory_targets
        if inventory_targets:
            inventory_matches = False
            if render_xml:
                sku_map = sku_ids_by_outer_id(render_xml)
                inventory_response = self.client.request("alibaba.icbu.product.sku.inventory.get", {
                    "language": "en_US", "product_id": product_id,
                })
                atomic_json(item_dir / "verify-inventory.json", inventory_response)
                rows = inventory_rows(inventory_response)
                target_rows = (
                    inventory_targets_from_xml(expected_xml)
                    if text(product.get("inventory_mode") or "deferred") == "embedded"
                    else [{
                        "sku_outer_id": target.get("sku_outer_id"),
                        "warehouse_code": target.get("warehouse_code"),
                        "stock_target": target.get("target"),
                    } for target in inventory_targets]
                )
                expected = {
                    (sku_map.get(text(target["sku_outer_id"])), text(target["warehouse_code"])): int(target["stock_target"])
                    for target in target_rows
                    if target.get("stock_target") is not None
                }
                actual = {
                    (text(row.get("sku_id")), text(row.get("inventory_code"))): int(row.get("inventory") or 0)
                    for row in rows
                }
                inventory_matches = bool(expected) and all(
                    sku_id and actual.get((sku_id, warehouse)) == target
                    for (sku_id, warehouse), target in expected.items()
                )
        checks = {
            "product_exists": bool(get_product),
            "status_approved": status == "approved",
            "display_yes": display == "Y",
            "category_matches": text(get_product.get("category_id")) == text(product["category_id"]),
            "title_matches": get_product.get("subject") == product["title"],
            "render_success": bool(render_payload.get("biz_success")),
            "changed_fields_match": fields_match,
            "inventory_matches": inventory_matches,
        }
        verified = all(checks.values())
        result = {
            "task_key": key,
            "product_id": product_id,
            "status": status,
            "display": display,
            "checks": checks,
            "verified": verified,
            "inventory_verified": inventory_matches,
            "pending": pending and not verified,
        }
        atomic_json(item_dir / "verification-report.json", result)
        return result

    def watch(self, max_wait_seconds: int) -> dict:
        ledger = WorkbookLedger(self.input_path, self.sheet)
        started = time.monotonic()
        delays = [0, 15, 30, 60, 120]
        cycle = 0
        final_results: dict[str, dict] = {}
        while True:
            pending_keys = [
                key for key, item in self.state["items"].items()
                if item.get("status") in {"auditing", "verification_pending"}
            ]
            if not pending_keys:
                break
            delay = delays[min(cycle, len(delays) - 1)]
            if delay:
                if time.monotonic() - started + delay > max_wait_seconds:
                    break
                time.sleep(delay)
            with ThreadPoolExecutor(max_workers=READ_CONCURRENCY) as executor:
                futures = {executor.submit(self.verify_item, key): key for key in pending_keys}
                for future in as_completed(futures):
                    key = futures[future]
                    try:
                        result = future.result()
                    except Exception as exc:
                        result = {
                            "task_key": key,
                            "verified": False,
                            "inventory_verified": False,
                            "pending": True,
                            "reason": f"{type(exc).__name__}: {exc}",
                        }
                    final_results[key] = result
                    item = self.state["items"][key]
                    if result.get("verified"):
                        item["status"] = "verified"
                        ledger.update(
                            item["sheet"], item["row"], "发布成功",
                            result["product_id"], "商品字段、展示状态及SKU库存验证通过",
                            time.strftime("%Y-%m-%dT%H:%M:%S%z"),
                        )
                    elif result.get("pending"):
                        item["status"] = "verification_pending"
                        ledger.update(
                            item["sheet"], item["row"], "待审核",
                            text(result.get("product_id")),
                            text(result.get("reason")) or "等待审核、字段或库存生效",
                        )
                    else:
                        item["status"] = "verification_failed"
                        ledger.update(
                            item["sheet"], item["row"], "验证失败",
                            text(result.get("product_id")),
                            text(result.get("reason")) or json.dumps(result.get("checks"), ensure_ascii=False),
                        )
            ledger.save()
            self.save_state()
            cycle += 1
            if time.monotonic() - started >= max_wait_seconds:
                break
        summary = {
            "verified": sum(item.get("status") == "verified" for item in self.state["items"].values()),
            "pending": sum(item.get("status") == "verification_pending" for item in self.state["items"].values()),
            "failed": sum(item.get("status") == "verification_failed" for item in self.state["items"].values()),
            "elapsed_seconds": round(time.monotonic() - started, 2),
        }
        atomic_json(self.root / "verification-summary.json", summary)
        return summary


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser()
    sub = root.add_subparsers(dest="command", required=True)
    for name in ("prepare", "submit", "reconcile-inventory", "watch", "run"):
        command = sub.add_parser(name)
        command.add_argument("--input", required=True, type=Path)
        command.add_argument("--work-dir", required=True, type=Path)
        command.add_argument("--sheet")
        command.add_argument("--config", "--mcp-config", dest="mcp_config", type=Path)
        if name == "prepare":
            command.add_argument("--upload-images", action="store_true")
            command.add_argument("--confirm", action="store_true")
        if name in {"submit", "reconcile-inventory", "run"}:
            command.add_argument("--confirm", action="store_true")
        if name == "watch":
            command.add_argument("--max-wait-seconds", type=int, default=900)
    return root


def main() -> None:
    args = parser().parse_args()
    if args.command == "prepare" and args.upload_images and not args.confirm:
        raise SystemExit("Photobank upload requires --confirm.")
    if args.command in {"submit", "reconcile-inventory", "run"} and not args.confirm:
        raise SystemExit("Formal write operations require --confirm.")
    workflow = Workflow(args.input, args.work_dir, args.sheet, args.mcp_config)
    if args.command == "prepare":
        result = workflow.prepare(args.upload_images)
    elif args.command == "submit":
        result = workflow.submit()
    elif args.command == "reconcile-inventory":
        result = workflow.reconcile_inventory()
    elif args.command == "watch":
        result = workflow.watch(args.max_wait_seconds)
    else:
        prepared = workflow.prepare(upload_images=True)
        submitted = workflow.submit()
        result = {"prepare": prepared, "submit": submitted}
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
