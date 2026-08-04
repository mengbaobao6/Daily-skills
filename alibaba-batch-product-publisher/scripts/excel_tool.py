#!/usr/bin/env python3
"""Inspect an Excel publishing ledger and perform verification-gated write-back."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("openpyxl is required to process Excel workbooks.") from exc


ALIASES = {
    "source_product_id": ("源商品ID", "source_product_id"),
    "title": ("新产品标题", "title"),
    "image_paths": ("新产品图片路径", "image_paths"),
    "specifications": ("新产品规格参数", "specifications"),
    "sku_plan_path": ("SKU方案文件路径", "sku_plan_path"),
    "task_id": ("任务ID", "task_id"),
}
OUTPUT_COLUMNS = ("发布状态", "新商品ID", "验证时间", "验证摘要")
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def natural_key(path: Path) -> list[Any]:
    return [int(part) if part.isdigit() else part.casefold()
            for part in re.split(r"(\d+)", path.name)]


def resolve_path(raw: str, workbook: Path) -> Path:
    path = Path(raw.strip().strip('"'))
    return path if path.is_absolute() else workbook.parent / path


def resolve_images(raw: str, workbook: Path) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    if not raw:
        return [], ["新产品图片路径为空"]
    folder = resolve_path(raw, workbook)
    if not folder.is_dir():
        return [], [f"新产品图片路径必须是存在的文件夹: {folder}"]
    images = sorted(
        (p.resolve() for p in folder.iterdir()
         if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS),
        key=natural_key,
    )
    if not images:
        errors.append(f"图片文件夹内没有支持的图片: {folder}")
    return [str(path) for path in images], errors


def find_columns(sheet) -> tuple[dict[str, int], list[str]]:
    headers = {text(cell.value): cell.column for cell in sheet[1] if text(cell.value)}
    columns: dict[str, int] = {}
    missing: list[str] = []
    for field, aliases in ALIASES.items():
        match = next((headers[name] for name in aliases if name in headers), None)
        if match:
            columns[field] = match
        elif field != "task_id":
            missing.append(aliases[0])
    return columns, missing


def inspect_workbook(input_path: Path, output_dir: Path, requested_sheet: str | None) -> None:
    keep_vba = input_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(input_path, data_only=False, keep_vba=keep_vba)
    sheets = [workbook[requested_sheet]] if requested_sheet else workbook.worksheets
    result: dict[str, Any] = {
        "workbook": str(input_path.resolve()),
        "sha256": hashlib.sha256(input_path.read_bytes()).hexdigest(),
        "sheets": [],
    }
    eligible = 0
    for sheet in sheets:
        columns, missing = find_columns(sheet)
        sheet_result: dict[str, Any] = {"sheet": sheet.title, "missing_columns": missing, "rows": []}
        if missing:
            result["sheets"].append(sheet_result)
            continue
        for row_number in range(2, sheet.max_row + 1):
            values = {name: text(sheet.cell(row_number, column).value)
                      for name, column in columns.items()}
            if not any(values.values()):
                continue
            errors: list[str] = []
            source_id = values["source_product_id"]
            if not source_id.isdigit():
                errors.append("源商品ID必须是数字")
            if not values["title"]:
                errors.append("新产品标题为空")
            images, image_errors = resolve_images(values["image_paths"], input_path)
            errors.extend(image_errors)
            sku_path = resolve_path(values["sku_plan_path"], input_path)
            if not sku_path.is_file():
                errors.append(f"SKU方案文件不存在: {sku_path}")
            specs = values["specifications"]
            spec_path = resolve_path(specs, input_path) if specs else None
            task_key = values.get("task_id") or f"{sheet.title}!{row_number}"
            row_result = {
                "task_key": task_key,
                "row": row_number,
                "source_product_id": source_id,
                "title": values["title"],
                "image_paths": images,
                "specifications": str(spec_path.resolve()) if spec_path and spec_path.is_file() else specs,
                "sku_plan_path": str(sku_path.resolve()),
                "eligible_for_platform_preflight": not errors,
                "errors": errors,
            }
            eligible += int(not errors)
            sheet_result["rows"].append(row_result)
        result["sheets"].append(sheet_result)
    result["eligible_rows"] = eligible
    output_dir.mkdir(parents=True, exist_ok=True)
    report = output_dir / f"{input_path.stem}-excel-inspection.json"
    report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"inspection": str(report), "eligible_rows": eligible}, ensure_ascii=False))


def verification_payload(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SystemExit(f"Cannot read verification JSON: {exc}") from exc
    if payload.get("verified") is not True or payload.get("inventory_verified") is not True:
        raise SystemExit("Refusing success write-back: product or inventory verification is incomplete.")
    return payload


def writeback(args: argparse.Namespace) -> None:
    payload = verification_payload(args.verification)
    if text(payload.get("source_product_id")) != args.source_product_id:
        raise SystemExit("Verification source product ID does not match the workbook row.")
    if text(payload.get("new_product_id")) != args.new_product_id:
        raise SystemExit("Verification new product ID does not match the requested write-back.")
    keep_vba = args.input.suffix.lower() == ".xlsm"
    workbook = load_workbook(args.input, data_only=False, keep_vba=keep_vba)
    if args.sheet not in workbook.sheetnames:
        raise SystemExit(f"Worksheet not found: {args.sheet}")
    sheet = workbook[args.sheet]
    columns, missing = find_columns(sheet)
    if missing:
        raise SystemExit(f"Workbook is missing required columns: {', '.join(missing)}")
    actual_source = text(sheet.cell(args.row, columns["source_product_id"]).value)
    actual_title = text(sheet.cell(args.row, columns["title"]).value)
    if actual_source != args.source_product_id or actual_title != args.title:
        raise SystemExit("Workbook row changed after preflight; refusing write-back.")
    headers = {text(cell.value): cell.column for cell in sheet[1] if text(cell.value)}
    for name in OUTPUT_COLUMNS:
        if name not in headers:
            column = sheet.max_column + 1
            sheet.cell(1, column, name)
            headers[name] = column
    existing_id = text(sheet.cell(args.row, headers["新商品ID"]).value)
    if existing_id and existing_id != args.new_product_id:
        raise SystemExit("Workbook row already contains a different new product ID.")
    timestamp = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
    summary = text(payload.get("summary")) or "商品字段、展示状态及SKU库存验证通过"
    sheet.cell(args.row, headers["发布状态"], "发布成功")
    sheet.cell(args.row, headers["新商品ID"], args.new_product_id)
    sheet.cell(args.row, headers["验证时间"], timestamp)
    sheet.cell(args.row, headers["验证摘要"], summary[:500])
    backup_dir = args.input.parent / ".publisher-backups"
    backup_dir.mkdir(exist_ok=True)
    backup = backup_dir / f"{args.input.stem}-{datetime.now():%Y%m%d-%H%M%S}{args.input.suffix}"
    shutil.copy2(args.input, backup)
    temporary = args.input.with_name(f".{args.input.stem}.publishing{args.input.suffix}")
    workbook.save(temporary)
    temporary.replace(args.input)
    print(json.dumps({"status": "发布成功", "workbook": str(args.input), "backup": str(backup)}, ensure_ascii=False))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    commands = parser.add_subparsers(dest="command", required=True)
    inspect_cmd = commands.add_parser("inspect")
    inspect_cmd.add_argument("--input", required=True, type=Path)
    inspect_cmd.add_argument("--output-dir", required=True, type=Path)
    inspect_cmd.add_argument("--sheet")
    write_cmd = commands.add_parser("writeback")
    write_cmd.add_argument("--input", required=True, type=Path)
    write_cmd.add_argument("--sheet", required=True)
    write_cmd.add_argument("--row", required=True, type=int)
    write_cmd.add_argument("--source-product-id", required=True)
    write_cmd.add_argument("--title", required=True)
    write_cmd.add_argument("--new-product-id", required=True)
    write_cmd.add_argument("--verification", required=True, type=Path)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    if args.command == "inspect":
        inspect_workbook(args.input, args.output_dir, args.sheet)
    else:
        writeback(args)


if __name__ == "__main__":
    main()
